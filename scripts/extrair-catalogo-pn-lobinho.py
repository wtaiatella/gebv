import os
import json
import openpyxl

SUBSTS_MAP = {
    62: 'Conquistar a seguinte insígnia no nível 1+: Insígnia do Aprender',
    71: 'Conquistar ao menos uma das seguintes especialidades no nível 1+: Empreendedorismo, Educação Financeira, Administração, Reparos Domésticos, Oratória',
    85: 'Conquistar ao menos uma das seguintes especialidades no nível 1+: Arte Digital, Artes Visuais, Artesanato, Comédia, Costura e Estilismo, Encadernação, Grafite, HQ, Maquete, Pintura, Propaganda e Marketing, Robótica, Videomaker',
    107: 'Conquistar ao menos uma das seguintes especialidades ou insígnias no nível 1+: Horticultura, Reduzir, Reciclar, Reutilizar',
    122: 'Conquistar ao menos uma das seguintes especialidades ou insígnias no nível 1+: Meteorologia, Escoteiros pela Energia Solar',
    144: 'Conquistar ao menos uma das seguintes especialidades ou insígnias no nível 1+: Ciências da Terra, Zoologia, Oceanologia, Botânica, Campeões da Natureza',
    163: 'Conquistar ao menos uma das seguintes especialidades no nível 1+: Acampamento, Excursões, Montanhismo, Pioneiria, Sobrevivência',
    175: 'Conquistar ao menos uma das seguintes especialidades ou insígnias no nível 1+: Defesa Civil, Mensageiros da Paz, Insígnia da Boa Ação',
    222: 'Conquistar ao menos uma das seguintes especialidades no nível 1+: Brasilidades, Genealogia, Informações Turísticas, Tradições dos Povos Indígenas',
    239: 'Conquistar ao menos uma das seguintes insígnias no nível 1+: Insígnia da Lusofonia, Insígnia do Cone-Sul, Diálogos pela Paz',
    250: 'Conquistar a seguinte especialidade no nível 1+: Escotismo Mundial',
    268: 'Conquistar ao menos uma das seguintes especialidades no nível 1+: Anatomia Humana, Prevenção em Saúde, Primeiros Socorros',
    275: 'Conquistar ao menos uma das seguintes especialidades ou insígnias no nível 1+: Yoga, Diálogo Inter-religioso',
    314: 'Conquistar ao menos uma das seguintes especialidades no nível 1+: Noções Desportivas, Nutrição',
    335: 'Conquistar a seguinte especialidade no nível 1+: Prevenção aos Vícios',
    349: 'Conquistar a seguinte especialidade no nível 1+: Prevenção ao Bullying'
}

# Mapeamento oficial dos 18 blocos para as atividades complementares de PA (Guia da UEB)
PA_BLOCOS_MAP = {
    'Aprendizagem Contínua e Desenvolvimento Vocacional': [
        'PTS-12', 'PTS-13', 'PTS-14', 'PTS-15', 'PTS-16', 'PTS-18', 'PTS-19', 'PTS-20',
        'PTS-28', 'PTS-29',
        'RC-23', 'RC-24', 'RC-25', 'RC-26', 'RC-27', 'RC-28', 'RC-29', 'RC-30'
    ],
    'Autonomia e Liderança': [
        'RC-16', 'RC-17'
    ],
    'Criatividade e Inovação': [
        'PTS-21', 'PTS-22', 'PTS-23'
    ],
    'Inteligência Emocional': [
        'PTS-31',
        'RC-18', 'RC-19', 'RC-20', 'RC-21', 'RC-22',
        'RC-31', 'RC-32', 'RC-37', 'RC-38', 'RC-46'
    ],
    'Consumo Responsável': [
        'PTS-49', 'PTS-50', 'PTS-51'
    ],
    'Mudanças Climáticas': [
        'RC-55'
    ],
    'Preservação da Biodiversidade': [
        'RC-52', 'RC-56'
    ],
    'Vida ao Ar Livre': [
        'PTS-10', 'PTS-17',
        'RC-9', 'RC-14', 'RC-15'
    ],
    'Comunidade': [
        'PTS-42',
        'RC-57', 'RC-58', 'RC-63', 'RC-64'
    ],
    'Democracia': [
        'PTS-41',
        'RC-51'
    ],
    'Herança Cultural': [
        'PTS-43', 'PTS-44', 'PTS-45',
        'RC-53'
    ],
    'Promoção da Paz': [
        'PTS-58', 'PTS-59',
        'RC-39', 'RC-40', 'RC-41', 'RC-42',
        'RC-59', 'RC-60', 'RC-61', 'RC-69', 'RC-70'
    ],
    'Valores': [
        'PTS-24', 'PTS-25', 'PTS-26', 'PTS-27', 'PTS-39', 'PTS-40',
        'RC-35', 'RC-36', 'RC-49', 'RC-50', 'RC-67', 'RC-68'
    ],
    'Cuidado com o Corpo': [
        'PTS-7', 'PTS-8',
        'RC-1', 'RC-2', 'RC-3', 'RC-6', 'RC-7', 'RC-8', 'RC-10', 'RC-11'
    ],
    'Espiritualidade': [
        'PTS-52', 'PTS-53', 'PTS-54', 'PTS-55', 'PTS-56', 'PTS-57',
        'RC-65', 'RC-66'
    ],
    'Hábitos Saudáveis': [
        'PTS-1', 'PTS-2', 'PTS-3', 'PTS-4', 'PTS-5', 'PTS-6', 'PTS-9', 'PTS-11',
        'RC-12', 'RC-13'
    ],
    'Saúde Mental': [
        'PTS-37', 'PTS-38',
        'RC-4', 'RC-5'
    ],
    'Vínculos Saudáveis': [
        'PTS-32', 'PTS-33', 'PTS-34', 'PTS-35', 'PTS-36', 'PTS-46', 'PTS-47', 'PTS-48',
        'RC-33', 'RC-34', 'RC-43', 'RC-44', 'RC-45', 'RC-46', 'RC-47', 'RC-48'
    ]
}

def main():
    excel_path = os.path.join(
        os.getcwd(),
        'docs',
        'transição',
        'alcateia',
        'Progressão Alcateia - Atualização do Programa - Modelo.xlsx'
    )
    
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f'Planilha não encontrada: {excel_path}')

    multiramo_path = os.path.join(os.getcwd(), 'data', 'catalogo', 'pa_catalogo_multiramo.json')
    if not os.path.exists(multiramo_path):
        raise FileNotFoundError(f'Catálogo multiramo PA não encontrado: {multiramo_path}')

    with open(multiramo_path, 'r', encoding='utf-8') as f:
        multiramo_data = json.load(f)

    lobinho_ativs = {a['identificacao']: a for a in multiramo_data.get('LOBINHO', {}).get('atividades', [])}
    print(f'Carregadas {len(lobinho_ativs)} atividades canônicas de Lobinho do PA.')

    print(f'Lendo planilha: {excel_path}')
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # 1. Metadados dos 18 blocos da aba Lobinho 1
    ws_l1 = wb['Lobinho 1']
    blocos_meta = {}
    for r in range(4, 22):
        eixo = str(ws_l1.cell(r, 1).value or '').strip()
        bloco = str(ws_l1.cell(r, 2).value or '').strip()
        fixa = int(ws_l1.cell(r, 3).value or 0)
        total_var = int(ws_l1.cell(r, 4).value or 0)
        qt_var = int(ws_l1.cell(r, 5).value or 0)
        subst = int(ws_l1.cell(r, 6).value or 0)
        total_realizar = int(ws_l1.cell(r, 7).value or 0)
        
        blocos_meta[bloco] = {
            'nr_ordem': r - 3,
            'eixo': eixo,
            'bloco': bloco,
            'nr_acoes_fixas_obrigatorias': fixa,
            'nr_acoes_variaveis_exigidas': qt_var,
            'total_variaveis_opcoes': total_var,
            'tem_substitutiva': subst > 0,
            'total_acoes_realizar': total_realizar
        }

    # 2. Leitura da aba Progressão Pessoal
    ws_pp = wb['Progressão Pessoal']
    eixos_set = []
    intencionalidades = {}
    
    # Agrupa linhas da planilha por bloco
    raw_rows_por_bloco = {}
    for r in range(3, ws_pp.max_row + 1):
        acao = ws_pp.cell(r, 5).value
        if not acao:
            continue
            
        eixo_str = str(ws_pp.cell(r, 1).value or '').strip()
        bloco_str = str(ws_pp.cell(r, 2).value or '').strip()
        inte_str = str(ws_pp.cell(r, 3).value or '').strip()
        tp_str = str(ws_pp.cell(r, 4).value or '').strip()
        acao_str = str(acao).strip()
        mod_str = str(ws_pp.cell(r, 6).value or '').strip() or 'Básico'
        req_str = str(ws_pp.cell(r, 7).value or '').strip()
        
        if eixo_str and eixo_str not in eixos_set:
            eixos_set.append(eixo_str)
            
        if bloco_str not in intencionalidades:
            intencionalidades[bloco_str] = inte_str
            
        if bloco_str not in raw_rows_por_bloco:
            raw_rows_por_bloco[bloco_str] = []

        raw_rows_por_bloco[bloco_str].append({
            'row': r,
            'eixo': eixo_str,
            'bloco': bloco_str,
            'tp': tp_str,
            'acao': acao_str,
            'modalidade': mod_str,
            'req': req_str
        })

    # 3. Montar blocos_list ordenados
    blocos_list = []
    for b_name, b_info in blocos_meta.items():
        blocos_list.append({
            'eixo': b_info['eixo'],
            'bloco': b_name,
            'nr_acoes_fixas_obrigatorias': b_info['nr_acoes_fixas_obrigatorias'],
            'nr_acoes_variaveis_exigidas': b_info['nr_acoes_variaveis_exigidas'],
            'nr_ordem': b_info['nr_ordem'],
            'ds_intencionalidade': intencionalidades.get(b_name, '')
        })
    blocos_list.sort(key=lambda x: x['nr_ordem'])

    # 4. Construção das ações e regras por bloco:
    # Ordem canônica em cada bloco: Fixas -> Variáveis Padronizadas -> PA Complementares -> Substitutivas
    acoes_list = []
    regras_list = []

    for b_info in blocos_list:
        b_name = b_info['bloco']
        eixo_name = b_info['eixo']
        rows_bloco = raw_rows_por_bloco.get(b_name, [])

        fixas_bloco = []
        variaveis_bloco = []
        substitutivas_bloco = []

        i = 0
        while i < len(rows_bloco):
            item = rows_bloco[i]
            acao_text = item['acao']
            
            # Agrupar sequências consecutivas de 'Conquistar Especialidade: ...'
            if acao_text.startswith('Conquistar Especialidade:'):
                esp_group = []
                while (
                    i < len(rows_bloco)
                    and rows_bloco[i]['acao'].startswith('Conquistar Especialidade:')
                ):
                    esp_group.append(rows_bloco[i])
                    i += 1
                    
                esp_nomes = []
                seen = set()
                for r_item in esp_group:
                    esp_name = r_item['acao'].replace('Conquistar Especialidade:', '').strip()
                    if esp_name and esp_name not in seen:
                        seen.add(esp_name)
                        esp_nomes.append(esp_name)
                        
                first_r = esp_group[0]['row']
                req_str = esp_group[0]['req'] or 'Opções variáveis'
                chave = f"{b_name}::{first_r}"
                
                ds_consolidada = (
                    f"Conquistar ao menos uma das seguintes especialidades no nível 1+: "
                    f"{', '.join(esp_nomes)}"
                )
                
                acao_obj = {
                    'chave': chave,
                    'excel_row': first_r,
                    'eixo': eixo_name,
                    'bloco': b_name,
                    'tp_acao': 'Variável',
                    'modalidade': 'Básico',
                    'ds_acao': ds_consolidada,
                    'regra_qtd_texto': req_str
                }
                
                regra_obj = {
                    'chave': chave,
                    'eixo': eixo_name,
                    'bloco': b_name,
                    'tipo_acao': 'Variável',
                    'excel_row': first_r,
                    'ds_acao': ds_consolidada,
                    'operacao': 'SEM_EQUIVALENCIA',
                    'descricao_origem': 'Sem equivalência mapeada',
                    'fl_requer_validacao_manual': True,
                    'detalhes_regra': {
                        'tipo': 'SEM_EQUIVALENCIA'
                    }
                }
                variaveis_bloco.append((acao_obj, regra_obj))
            else:
                r = item['row']
                if r in SUBSTS_MAP:
                    acao_text = SUBSTS_MAP[r]
                chave = f"{b_name}::{r}"
                tp_str = item['tp']
                
                tipo_normalizado = 'Variável'
                if tp_str.lower().startswith('fix'):
                    tipo_normalizado = 'Fixa'
                elif 'substitu' in tp_str.lower():
                    tipo_normalizado = 'Substitutiva'
                    
                mod_normalizada = 'Básico'
                if item['modalidade'].lower() == 'ar':
                    mod_normalizada = 'Ar'
                elif item['modalidade'].lower() == 'mar':
                    mod_normalizada = 'Mar'
                    
                acao_obj = {
                    'chave': chave,
                    'excel_row': r,
                    'eixo': eixo_name,
                    'bloco': b_name,
                    'tp_acao': tipo_normalizado,
                    'modalidade': mod_normalizada,
                    'ds_acao': acao_text,
                    'regra_qtd_texto': item['req']
                }
                
                regra_obj = {
                    'chave': chave,
                    'eixo': eixo_name,
                    'bloco': b_name,
                    'tipo_acao': tipo_normalizado,
                    'excel_row': r,
                    'ds_acao': acao_text,
                    'operacao': 'SEM_EQUIVALENCIA',
                    'descricao_origem': 'Sem equivalência mapeada',
                    'fl_requer_validacao_manual': True,
                    'detalhes_regra': {
                        'tipo': 'SEM_EQUIVALENCIA'
                    }
                }
                
                if tipo_normalizado == 'Fixa':
                    fixas_bloco.append((acao_obj, regra_obj))
                elif tipo_normalizado == 'Substitutiva':
                    substitutivas_bloco.append((acao_obj, regra_obj))
                else:
                    variaveis_bloco.append((acao_obj, regra_obj))
                i += 1

        # 4.1 Injeta ações Fixas
        for ac, rg in fixas_bloco:
            acoes_list.append(ac)
            regras_list.append(rg)

        # 4.2 Injeta ações Variáveis Padrão
        for ac, rg in variaveis_bloco:
            acoes_list.append(ac)
            regras_list.append(rg)

        # 4.3 Injeta ações de PA Complementares do Bloco
        pa_items_bloco = PA_BLOCOS_MAP.get(b_name, [])
        for ident in pa_items_bloco:
            ativ = lobinho_ativs.get(ident)
            if not ativ:
                print(f'ALERTA: Atividade {ident} não encontrada no catálogo multiramo!')
                continue
            ds_texto = ativ['ds_atividade']
            chave = f"{b_name}::{ident}"

            acao_pa = {
                'chave': chave,
                'excel_row': 0,
                'eixo': eixo_name,
                'bloco': b_name,
                'tp_acao': 'Variável',
                'modalidade': 'PA',
                'ds_acao': ds_texto,
                'regra_qtd_texto': f"PA Complementar ({ident})"
            }

            regra_pa = {
                'chave': chave,
                'eixo': eixo_name,
                'bloco': b_name,
                'tipo_acao': 'Variável',
                'excel_row': 0,
                'ds_acao': ds_texto,
                'operacao': 'PROGRESSOES',
                'descricao_origem': ident,
                'fl_requer_validacao_manual': False,
                'detalhes_regra': {
                    'tipo': 'PROGRESSOES',
                    'item': {
                        'pa_atividade_id': 0,
                        'identificacao': ident,
                        'ds_atividade': ds_texto
                    }
                }
            }

            acoes_list.append(acao_pa)
            regras_list.append(regra_pa)

        # 4.4 Injeta ações Substitutivas
        for ac, rg in substitutivas_bloco:
            acoes_list.append(ac)
            regras_list.append(rg)

    # Estruturas finais
    catalogo_pn_lobinho = {
        'ramo': 'LOBINHO',
        'eixos': eixos_set,
        'blocos': blocos_list,
        'acoes': acoes_list
    }
    
    regras_pn_lobinho = {
        'versao': '2.0.0-lobinho',
        'ramo': 'LOBINHO',
        'total_regras': len(regras_list),
        'estatisticas': {
            'total': len(regras_list),
            'sem_equivalencia': len([r for r in regras_list if r['operacao'] == 'SEM_EQUIVALENCIA']),
            'progressoes': len([r for r in regras_list if r['operacao'] == 'PROGRESSOES']),
            'fixas': len([a for a in acoes_list if a['tp_acao'] == 'Fixa']),
            'variaveis': len([a for a in acoes_list if a['tp_acao'] == 'Variável' and a['modalidade'] != 'PA']),
            'pa_complementares': len([a for a in acoes_list if a['modalidade'] == 'PA']),
            'substitutivas': len([a for a in acoes_list if a['tp_acao'] == 'Substitutiva'])
        },
        'regras': regras_list
    }

    out_dir = os.path.join(os.getcwd(), 'data', 'catalogo')
    os.makedirs(out_dir, exist_ok=True)
    
    cat_path = os.path.join(out_dir, 'pn_catalogo_lobinho.json')
    req_path = os.path.join(out_dir, 'pn_equivalencia_regras_lobinho.json')
    
    with open(cat_path, 'w', encoding='utf-8') as f:
        json.dump(catalogo_pn_lobinho, f, indent=2, ensure_ascii=False)
        
    with open(req_path, 'w', encoding='utf-8') as f:
        json.dump(regras_pn_lobinho, f, indent=2, ensure_ascii=False)

    print(f'\n✓ Catálogo gerado: {cat_path}')
    print(f'  - Eixos: {len(eixos_set)}')
    print(f'  - Blocos: {len(blocos_list)}')
    print(f'  - Ações totais: {len(acoes_list)}')
    print(f'    • Fixas: {regras_pn_lobinho["estatisticas"]["fixas"]}')
    print(f'    • Variáveis padrão: {regras_pn_lobinho["estatisticas"]["variaveis"]}')
    print(f'    • PA complementares: {regras_pn_lobinho["estatisticas"]["pa_complementares"]}')
    print(f'    • Substitutivas: {regras_pn_lobinho["estatisticas"]["substitutivas"]}')
    print(f'✓ Regras geradas: {req_path}')
    print(f'  - Total de regras: {len(regras_list)}')
    print(f'    • PROGRESSOES (PA): {regras_pn_lobinho["estatisticas"]["progressoes"]}')
    print(f'    • SEM_EQUIVALENCIA: {regras_pn_lobinho["estatisticas"]["sem_equivalencia"]}')

if __name__ == '__main__':
    main()
