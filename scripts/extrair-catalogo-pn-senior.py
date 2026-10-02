#!/usr/bin/env python3
"""
Script de extração do Catálogo do Novo Programa (PN) e Regras de Equivalência do Ramo Sênior
a partir da planilha oficial SENIOR.xlsx e catálogo multiramo de PA.

Gera:
  - data/catalogo/pn_catalogo_senior.json
  - data/catalogo/pn_equivalencia_regras_senior.json
"""

import os
import json
import re
import openpyxl

# Mapeamento oficial dos 18 blocos para as atividades complementares de PA (Guia da UEB / Coluna G)
PA_BLOCOS_MAP = {
    # Eixo: Habilidades para a Vida
    'Aprendizagem Contínua e Desenvolvimento Vocacional': [
        'ECA-33', 'ECA-35', 'ECA-38', 'ECA-41', 'ECA-42', 'ECA-45', 'ECA-46'
    ],
    'Autonomia e Liderança': [
        'ECA-33', 'ECA-34', 'ECA-35', 'ECA-36', 'ECA-37', 'ECA-38'
    ],
    'Criatividade e Inovação': [
        'ECA-36'
    ],
    'Inteligência Emocional': [
        'ECA-47', 'ECA-48', 'ECA-49'
    ],

    # Eixo: Meio Ambiente
    'Consumo Responsável': [],
    'Mudanças Climáticas': [
        'ECA-66', 'ECA-67'
    ],
    'Preservação da Biodiversidade': [
        'ECA-68', 'ECA-69'
    ],
    'Vida ao Ar Livre': [
        'ECA-21', 'ECA-22', 'ECA-23', 'ECA-24', 'ECA-25', 'ECA-26',
        'ECA-27', 'ECA-28', 'ECA-29', 'ECA-30', 'ECA-31', 'ECA-32'
    ],

    # Eixo: Paz e Desenvolvimento
    'Comunidade': [
        'ECA-51', 'ECA-52', 'ECA-53', 'ECA-61'
    ],
    'Democracia': [
        'ECA-58', 'ECA-59', 'ECA-60'
    ],
    'Herança Cultural': [
        'ECA-62', 'ECA-63'
    ],
    'Promoção da Paz': [
        'ECA-57', 'ECA-64', 'ECA-65', 'ECA-74', 'ECA-75'
    ],
    'Valores': [
        'ECA-39', 'ECA-40', 'ECA-43'
    ],

    # Eixo: Saúde e Bem-Estar
    'Cuidado com o Corpo': [
        'ECA-1', 'ECA-2', 'ECA-3', 'ECA-4', 'ECA-5', 'ECA-6',
        'ECA-7', 'ECA-8', 'ECA-9', 'ECA-10'
    ],
    'Espiritualidade': [
        'ECA-70', 'ECA-71', 'ECA-72', 'ECA-73'
    ],
    'Hábitos Saudáveis': [
        'ECA-11', 'ECA-12', 'ECA-13', 'ECA-14', 'ECA-15', 'ECA-16',
        'ECA-17', 'ECA-18', 'ECA-19'
    ],
    'Saúde Mental': [
        'ECA-20', 'ECA-37', 'ECA-44'
    ],
    'Vínculos Saudáveis': [
        'ECA-50', 'ECA-54', 'ECA-55', 'ECA-56'
    ]
}

EIXO_NAMES_MAP = {
    'HABILIDADES PARA VIDA': 'Habilidades para a Vida',
    'MEIO AMBIENTE': 'Meio Ambiente',
    'PAZ E DESENVOLVIMENTO': 'Paz e Desenvolvimento',
    'SAUDE E BEM-ESTAR': 'Saúde e Bem-Estar'
}

TEXTO_PARA_NUMERO = {
    'uma': 1, 'duas': 2, 'tres': 3, 'três': 3, 'quatro': 4,
    'cinco': 5, 'seis': 6, 'sete': 7, 'oito': 8
}

def extrair_qtd_exigida(texto: str) -> int:
    m = re.search(r'realizar\s+([a-zá-ú]+)\s+ações', texto.lower())
    if m:
        palavra = m.group(1).strip()
        return TEXTO_PARA_NUMERO.get(palavra, 2)
    return 2

def formatar_substitutiva(c3_text: str, prev_c3_text: str) -> str:
    cleaned = c3_text.strip().rstrip('.')
    is_insignia = ('insígne' in prev_c3_text.lower() or 
                   'insígnia' in prev_c3_text.lower() or 
                   'insígnia' in cleaned.lower())

    if 'reduzir' in cleaned.lower() and 'reciclar' in cleaned.lower():
        return 'Conquistar a seguinte insígnia no nível 1+: Reduzir, Reciclar, Reutilizar'

    if is_insignia:
        items = [p.strip().rstrip('.') for p in re.split(r'[;]', cleaned) if p.strip()]
        if len(items) <= 1:
            items = [p.strip().rstrip('.') for p in re.split(r'[,]', cleaned) if p.strip()]
        if len(items) == 1:
            return f'Conquistar a seguinte insígnia no nível 1+: {items[0]}'
        else:
            return f'Conquistar ao menos uma das seguintes insígnias no nível 1+: {", ".join(items)}'
    else:
        items = [p.strip().rstrip('.') for p in re.split(r'[;,]', cleaned) if p.strip()]
        if len(items) == 1:
            return f'Conquistar a seguinte especialidade no nível 1+: {items[0]}'
        else:
            return f'Conquistar ao menos uma das seguintes especialidades no nível 1+: {", ".join(items)}'

def main():
    excel_path = os.path.join(
        os.getcwd(),
        'docs',
        'transição',
        'senior',
        'SENIOR.xlsx'
    )
    if not os.path.exists(excel_path):
        excel_path = os.path.join(
            os.path.dirname(os.getcwd()),
            'gebv',
            'docs',
            'transição',
            'senior',
            'SENIOR.xlsx'
        )
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f'Planilha não encontrada: {excel_path}')

    multiramo_path = os.path.join(os.getcwd(), 'data', 'catalogo', 'pa_catalogo_multiramo.json')
    if not os.path.exists(multiramo_path):
        multiramo_path = os.path.join(os.path.dirname(os.getcwd()), 'gebv', 'data', 'catalogo', 'pa_catalogo_multiramo.json')
    if not os.path.exists(multiramo_path):
        raise FileNotFoundError(f'Catálogo multiramo PA não encontrado: {multiramo_path}')

    with open(multiramo_path, 'r', encoding='utf-8') as f:
        multiramo_data = json.load(f)

    senior_ativs = {a['identificacao']: a for a in multiramo_data.get('SENIOR', {}).get('atividades', [])}
    print(f'Carregadas {len(senior_ativs)} atividades canônicas de Sênior do PA.')

    print(f'Lendo planilha: {excel_path}')
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    eixos_lista = [
        'Habilidades para a Vida',
        'Meio Ambiente',
        'Paz e Desenvolvimento',
        'Saúde e Bem-Estar'
    ]
    blocos_dict = {}
    acoes_list = []
    regras_list = []

    sheets_to_process = ['HABILIDADES PARA VIDA', 'MEIO AMBIENTE', 'PAZ E DESENVOLVIMENTO', 'SAUDE E BEM-ESTAR']

    bloco_order = 1
    total_fixas = 0
    total_variaveis = 0
    total_substs = 0
    total_pa = 0

    for sheet_name in sheets_to_process:
        if sheet_name not in wb.sheetnames:
            print(f'Aviso: Aba {sheet_name} não encontrada!')
            continue

        ws = wb[sheet_name]
        eixo_name = EIXO_NAMES_MAP.get(sheet_name, sheet_name)

        # 1. Agrupa linhas por bloco
        blocos_linhas = []
        bloco_atual = None

        for r in range(1, ws.max_row + 1):
            c1 = ws.cell(r, 1).value
            c2 = ws.cell(r, 2).value
            c3 = ws.cell(r, 3).value
            c4 = ws.cell(r, 4).value
            c7 = ws.cell(r, 7).value

            if c1 and str(c1).strip() != 'Eixo':
                b_name = str(c1).strip()
                b_intent = str(c2).strip() if c2 else ''
                bloco_atual = {
                    'nome': b_name,
                    'intencionalidade': b_intent,
                    'sheet': sheet_name,
                    'eixo': eixo_name,
                    'rows': []
                }
                blocos_linhas.append(bloco_atual)

            if bloco_atual and r > 1:
                bloco_atual['rows'].append({
                    'row': r,
                    'c1': c1,
                    'c2': c2,
                    'c3': str(c3).strip() if c3 else None,
                    'c4': str(c4).strip() if c4 else None,
                    'c7': str(c7).strip() if c7 else None
                })

        # 2. Processa cada bloco da aba
        for b_info in blocos_linhas:
            b_name = b_info['nome']
            b_intent = b_info['intencionalidade']

            fixas_bloco = []
            variaveis_bloco = []
            substitutivas_bloco = []
            qtd_vars_exigida = 2

            for idx, r_data in enumerate(b_info['rows']):
                c3 = r_data['c3']
                c4 = r_data['c4']
                row_idx = r_data['row']

                if not c3:
                    continue

                if 'dentre as listadas' in c3.lower() or 'ações dentre' in c3.lower():
                    qtd_vars_exigida = extrair_qtd_exigida(c3)

                if c4 == 'Fixa':
                    chave = f"{b_name}::{row_idx}"
                    acao_obj = {
                        'chave': chave,
                        'excel_row': row_idx,
                        'eixo': eixo_name,
                        'bloco': b_name,
                        'tp_acao': 'Fixa',
                        'modalidade': 'Básico',
                        'ds_acao': c3,
                        'regra_qtd_texto': 'Ação Fixa'
                    }
                    regra_obj = {
                        'chave': chave,
                        'eixo': eixo_name,
                        'bloco': b_name,
                        'tipo_acao': 'Fixa',
                        'excel_row': row_idx,
                        'ds_acao': c3,
                        'operacao': 'SEM_EQUIVALENCIA',
                        'descricao_origem': 'Sem equivalência mapeada',
                        'fl_requer_validacao_manual': True,
                        'detalhes_regra': {'tipo': 'SEM_EQUIVALENCIA'}
                    }
                    fixas_bloco.append((acao_obj, regra_obj))

                elif c4 == 'Variável':
                    chave = f"{b_name}::{row_idx}"
                    acao_obj = {
                        'chave': chave,
                        'excel_row': row_idx,
                        'eixo': eixo_name,
                        'bloco': b_name,
                        'tp_acao': 'Variável',
                        'modalidade': 'Básico',
                        'ds_acao': c3,
                        'regra_qtd_texto': 'Ação Variável'
                    }
                    regra_obj = {
                        'chave': chave,
                        'eixo': eixo_name,
                        'bloco': b_name,
                        'tipo_acao': 'Variável',
                        'excel_row': row_idx,
                        'ds_acao': c3,
                        'operacao': 'SEM_EQUIVALENCIA',
                        'descricao_origem': 'Sem equivalência mapeada',
                        'fl_requer_validacao_manual': True,
                        'detalhes_regra': {'tipo': 'SEM_EQUIVALENCIA'}
                    }
                    variaveis_bloco.append((acao_obj, regra_obj))

                elif c4 == 'Substiuir Variavel':
                    prev_c3 = b_info['rows'][idx - 1]['c3'] if idx > 0 else ''
                    ds_subst = formatar_substitutiva(c3, prev_c3 or '')
                    chave = f"{b_name}::{row_idx}"
                    acao_obj = {
                        'chave': chave,
                        'excel_row': row_idx,
                        'eixo': eixo_name,
                        'bloco': b_name,
                        'tp_acao': 'Substitutiva',
                        'modalidade': 'Básico',
                        'ds_acao': ds_subst,
                        'regra_qtd_texto': 'Ação Substitutiva'
                    }
                    regra_obj = {
                        'chave': chave,
                        'eixo': eixo_name,
                        'bloco': b_name,
                        'tipo_acao': 'Substitutiva',
                        'excel_row': row_idx,
                        'ds_acao': ds_subst,
                        'operacao': 'SEM_EQUIVALENCIA',
                        'descricao_origem': 'Sem equivalência mapeada',
                        'fl_requer_validacao_manual': True,
                        'detalhes_regra': {'tipo': 'SEM_EQUIVALENCIA'}
                    }
                    substitutivas_bloco.append((acao_obj, regra_obj))

            if b_name not in blocos_dict:
                blocos_dict[b_name] = {
                    'eixo': eixo_name,
                    'bloco': b_name,
                    'ds_intencionalidade': b_intent,
                    'nr_acoes_fixas_obrigatorias': len(fixas_bloco),
                    'nr_acoes_variaveis_exigidas': qtd_vars_exigida,
                    'nr_ordem': bloco_order
                }
                bloco_order += 1

            # Injeta Fixas
            for ac, rg in fixas_bloco:
                acoes_list.append(ac)
                regras_list.append(rg)
                total_fixas += 1

            # Injeta Variáveis Padrão
            for ac, rg in variaveis_bloco:
                acoes_list.append(ac)
                regras_list.append(rg)
                total_variaveis += 1

            # Injeta Ações de PA Complementares do Bloco
            pa_items_bloco = PA_BLOCOS_MAP.get(b_name, [])
            for ident in pa_items_bloco:
                ativ = senior_ativs.get(ident)
                if not ativ:
                    print(f'ALERTA: Atividade {ident} não encontrada no catálogo multiramo!')
                    continue
                ds_texto = ativ['ds_atividade']
                chave = f"{b_name}::{ident}"

                acao_pa = {
                    'chave': chave,
                    'excel_row': 0,
                    'eixo': eixo_name,
                    'bloco': b_name,
                    'tp_acao': 'Variável',
                    'modalidade': 'PA',
                    'ds_acao': ds_texto,
                    'regra_qtd_texto': f'Atividade PA {ident}'
                }

                regra_pa = {
                    'chave': chave,
                    'eixo': eixo_name,
                    'bloco': b_name,
                    'tipo_acao': 'Variável',
                    'excel_row': 0,
                    'ds_acao': ds_texto,
                    'operacao': 'PROGRESSOES',
                    'descricao_origem': ident,
                    'fl_requer_validacao_manual': False,
                    'detalhes_regra': {
                        'tipo': 'PROGRESSOES',
                        'item': {
                            'pa_atividade_id': 0,
                            'identificacao': ident,
                            'ds_atividade': ds_texto
                        }
                    }
                }

                acoes_list.append(acao_pa)
                regras_list.append(regra_pa)
                total_pa += 1

            # Injeta Substitutivas
            for ac, rg in substitutivas_bloco:
                acoes_list.append(ac)
                regras_list.append(rg)
                total_substs += 1

    print('\n========================================')
    print('          RESUMO DA EXTRAÇÃO SÊNIOR     ')
    print('========================================')
    print(f'Eixos processados:         {len(eixos_lista)}')
    print(f'Blocos processados:        {len(blocos_dict)}')
    print(f'Ações Fixas:               {total_fixas}')
    print(f'Ações Variáveis (PN):      {total_variaveis}')
    print(f'Ações PA Complementares:   {total_pa}')
    print(f'Ações Substitutivas:       {total_substs}')
    print(f'TOTAL DE AÇÕES EDUCATIVAS: {len(acoes_list)}')
    print(f'TOTAL DE REGRAS:           {len(regras_list)}')
    print('========================================')

    out_dir = os.path.join(os.getcwd(), 'data', 'catalogo')
    os.makedirs(out_dir, exist_ok=True)

    catalogo_pn = {
        'ramo': 'SENIOR',
        'eixos': eixos_lista,
        'blocos': list(blocos_dict.values()),
        'acoes': acoes_list
    }

    pn_cat_path = os.path.join(out_dir, 'pn_catalogo_senior.json')
    with open(pn_cat_path, 'w', encoding='utf-8') as f:
        json.dump(catalogo_pn, f, indent=2, ensure_ascii=False)
    print(f'✓ Salvo: {pn_cat_path}')

    equivalencias_json = {
        'versao': '2.0.0-senior',
        'ramo': 'SENIOR',
        'total_regras': len(regras_list),
        'estatisticas': {
            'total': len(regras_list),
            'sem_equivalencia': len(regras_list) - total_pa,
            'progressoes': total_pa,
            'fixas': total_fixas,
            'variaveis': total_variaveis,
            'pa_complementares': total_pa,
            'substitutivas': total_substs
        },
        'regras': regras_list
    }

    regras_cat_path = os.path.join(out_dir, 'pn_equivalencia_regras_senior.json')
    with open(regras_cat_path, 'w', encoding='utf-8') as f:
        json.dump(equivalencias_json, f, indent=2, ensure_ascii=False)
    print(f'✓ Salvo: {regras_cat_path}')

if __name__ == '__main__':
    main()
