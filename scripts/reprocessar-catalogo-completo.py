#!/usr/bin/env python3
"""
Reprocessador Oficial do Catálogo do Novo Programa e Matriz de Equivalências
Lê diretamente as abas da planilha com as classificações na Coluna B e regras na Coluna D+:
  - docs/transição/escoteiro/2026 02 26 - Equiparação progressão ramo escoteiro.xlsx
Gera:
  - data/catalogo/pn_catalogo.json
  - data/catalogo/equivalencias.json
"""

import os
import json
import re
import unicodedata
import openpyxl

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))

def resolve_doc_path(rel_subpath):
    p_trans = os.path.join(BASE_DIR, 'docs/transição/escoteiro', rel_subpath)
    p_norm = os.path.join(BASE_DIR, 'docs/escoteiro', rel_subpath)
    if os.path.exists(p_trans): return p_trans
    return p_norm

EXCEL_PATH = resolve_doc_path('2026 02 26 - Equiparação progressão ramo escoteiro.xlsx')
PA_CATALOGO_PATH = os.path.join(BASE_DIR, 'data/catalogo/pa_catalogo.json')
ESP_CATALOGO_PATH = os.path.join(BASE_DIR, 'data/pa_especialidades_catalogo.json')

with open(PA_CATALOGO_PATH, 'r', encoding='utf-8') as f:
    pa_catalogo = json.load(f)

with open(ESP_CATALOGO_PATH, 'r', encoding='utf-8') as f:
    esp_catalogo = json.load(f)

def normalize_text_key(text):
    if not text: return ''
    s = unicodedata.normalize('NFD', str(text))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return ''.join(c for c in s.lower() if c.isalnum())

pa_atividades_map = {}
if 'pistas_items' in pa_catalogo:
    for item in pa_catalogo['pistas_items']:
        cd_ueb = str(item.get('cd_ueb', '')).strip()
        nr_num = re.sub(r'\D', '', cd_ueb)
        if nr_num:
            pa_atividades_map[f"PT-{nr_num}"] = {
                'cd_ueb': nr_num,
                'cd_caminho_paxtu': '5',
                'ds_competencia': item.get('ds_competencia', ''),
                'ds_area': item.get('ds_area', ''),
                'ds_atividade': item.get('ds_atividade', '')
            }

if 'rumo_items' in pa_catalogo:
    for item in pa_catalogo['rumo_items']:
        cd_ueb = str(item.get('cd_ueb', '')).strip()
        nr_num = re.sub(r'\D', '', cd_ueb)
        if nr_num:
            pa_atividades_map[f"RT-{nr_num}"] = {
                'cd_ueb': nr_num,
                'cd_caminho_paxtu': '6',
                'ds_competencia': item.get('ds_competencia', ''),
                'ds_area': item.get('ds_area', ''),
                'ds_atividade': item.get('ds_atividade', '')
            }

def clean_str(s):
    if s is None:
        return ''
    s = str(s).replace('_x000d_', '').replace('\r', '')
    return s.strip()

def parse_num_words(text):
    text_lower = text.lower()
    if 'uma' in text_lower or 'um' in text_lower: return 1
    if 'duas' in text_lower or 'dois' in text_lower: return 2
    if 'tres' in text_lower or 'três' in text_lower: return 3
    if 'quatro' in text_lower: return 4
    if 'cinco' in text_lower: return 5
    if 'seis' in text_lower: return 6
    if 'sete' in text_lower: return 7
    if 'oito' in text_lower: return 8
    nums = re.findall(r'\d+', text)
    if nums:
        return int(nums[0])
    return 1

def parse_pa_refs_from_text(raw_text):
    """Extrai refs de Pistas (PT) e Rumo (RT) de textos como:
    Pista 37, Pista 82 e 83, Pistas 21 a 24, Rumo 42, P55 a 58, rumo 12 a 14
    """
    pistas = []
    rumo = []
    text = clean_str(raw_text)
    
    # Intervalos de Pista
    for m in re.finditer(r'(?:pistas?|p)\s*(\d+)\s*(?:a|à|-|e)\s*(\d+)', text, re.IGNORECASE):
        start_n = int(m.group(1))
        end_n = int(m.group(2))
        for num in range(start_n, end_n + 1):
            pistas.append(str(num))
            
    # Intervalos de Rumo
    for m in re.finditer(r'(?:rumo|r)\s*(\d+)\s*(?:a|à|-|e)\s*(\d+)', text, re.IGNORECASE):
        start_n = int(m.group(1))
        end_n = int(m.group(2))
        for num in range(start_n, end_n + 1):
            rumo.append(str(num))

    # Pistas avulsas
    for m in re.finditer(r'\b(?:pista|pistas|p)\s*(\d+)\b', text, re.IGNORECASE):
        num = m.group(1)
        if num not in pistas:
            pistas.append(num)
            
    # Rumos avulsos
    for m in re.finditer(r'\b(?:rumo|rumos|r)\s*(\d+)\b', text, re.IGNORECASE):
        num = m.group(1)
        if num not in rumo:
            rumo.append(num)

    # Remove duplicados preservando ordem
    pistas_clean = []
    for p in pistas:
        if p not in pistas_clean: pistas_clean.append(p)
    rumo_clean = []
    for r in rumo:
        if r not in rumo_clean: rumo_clean.append(r)

    return pistas_clean, rumo_clean

def parse_especialidades_from_text(raw_text):
    """Extrai nomes de especialidades e nível mínimo de strings como:
    Aeronáutica N1+, Administração N2+, Internet N1+, Yoga
    """
    esps = []
    text = clean_str(raw_text)
    if not text:
        return esps, 1
    
    nivel = 1
    if 'n2+' in text.lower() or 'nível 2' in text.lower() or 'nivel 2' in text.lower() or 'n2' in text.lower():
        nivel = 2
    elif 'n3+' in text.lower() or 'n3' in text.lower():
        nivel = 3
        
    # Limpa sufixos de nível
    cleaned = re.sub(r'\s*(?:N[123]\+|n[123]\+|N[123]|n[123]|nível\s*[123]|nivel\s*[123])', '', text, flags=re.IGNORECASE).strip()
    cleaned = re.sub(r'^(?:especialidade\s*|insígnia\s*|insignia\s*)', '', cleaned, flags=re.IGNORECASE).strip()
    
    if cleaned and not re.match(r'^(?:pistas?|rumo|\d+)$', cleaned, re.IGNORECASE):
        parts = re.split(r'[,;]|\bOR\b|\be\b', cleaned, flags=re.IGNORECASE)
        for p in parts:
            p_str = p.strip()
            if p_str and len(p_str) > 2 and not p_str.lower().startswith('rumo') and not p_str.lower().startswith('pista'):
                esps.append(p_str)
                
    return esps, nivel

def parse_atividades_pa_bloco(raw_text):
    """Extrai a lista de itens PA complementares do cabeçalho do bloco (ex: I22, I23, F4, etc.)"""
    pt_items = []
    rt_items = []
    
    lines = clean_str(raw_text).split('\n')
    current_section = None
    
    for l in lines:
        l_str = l.strip()
        if not l_str:
            continue
        if 'pistas e trilha' in l_str.lower():
            current_section = 'PT'
            continue
        elif 'rumo e travessia' in l_str.lower():
            current_section = 'RT'
            continue
        
        tokens = re.findall(r'\b[A-Za-z]?(\d+)\b', l_str)
        if current_section == 'PT':
            for t in tokens:
                if t not in pt_items:
                    pt_items.append(t)
        elif current_section == 'RT':
            for t in tokens:
                if t not in rt_items:
                    rt_items.append(t)
                    
    return pt_items, rt_items

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

target_sheets = [
    'Habilidades para a Vida',
    'Meio Ambiente',
    'Paz e Desenvolvimento',
    'Saúde e Bem Estar'
]

eixos_list = [
    'Habilidades para a Vida',
    'Meio Ambiente',
    'Paz e Desenvolvimento',
    'Saúde e Bem-Estar'
]

parsed_blocos = []
parsed_acoes = []
parsed_regras = []
bloco_global_ordem = 0

for sheet_name in target_sheets:
    ws = wb[sheet_name]
    eixo_nm = 'Saúde e Bem-Estar' if 'Saúde' in sheet_name or 'Saude' in sheet_name else sheet_name
    print(f"\n=======================================================")
    print(f"PROCESSANDO EIXO: {eixo_nm} ({ws.max_row} linhas)")
    print(f"=======================================================")
    
    # Agrupar linhas por Bloco
    blocos_raw = []
    current_bloco_rows = []
    
    for r in range(1, ws.max_row + 1):
        c1 = clean_str(ws.cell(r, 1).value)
        c2 = clean_str(ws.cell(r, 2).value)
        c3 = clean_str(ws.cell(r, 3).value)
        row_vals = [ws.cell(r, c).value for c in range(4, 25)]
        while row_vals and row_vals[-1] is None:
            row_vals.pop()
        row_vals_clean = [clean_str(x) for x in row_vals if clean_str(x)]
        
        if 'inicio do bloco' in c2.lower() or 'inicio do bloco' in c1.lower():
            if current_bloco_rows:
                blocos_raw.append(current_bloco_rows)
            current_bloco_rows = [{'row': r, 'tipo': c2, 'texto': c3 if c3 else c2, 'regras': row_vals_clean}]
        else:
            if current_bloco_rows and (c2 or c3 or row_vals_clean):
                current_bloco_rows.append({'row': r, 'tipo': c2, 'texto': c3, 'regras': row_vals_clean})
                
    if current_bloco_rows:
        blocos_raw.append(current_bloco_rows)
        
    print(f"Total de blocos detectados em {sheet_name}: {len(blocos_raw)}")
    
    for b_idx, b_rows in enumerate(blocos_raw):
        bloco_global_ordem += 1
        nm_bloco = b_rows[0]['texto']
        print(f"\n--- Bloco {bloco_global_ordem}: {nm_bloco} ---")
        
        intencionalidade = ""
        pa_complementares_raw = ""
        variaveis_exigidas = 4
        
        fixas_list = []
        variaveis_list = []
        substitutivas_list = []
        esps_compl = []
        
        for item in b_rows:
            tipo = item['tipo'].lower()
            texto = item['texto']
            regras = item['regras']
            r_idx = item['row']
            
            if not tipo and not texto and not regras:
                continue
                
            if 'intencionalidade' in tipo or 'intencionalidade' in texto.lower():
                intencionalidade = texto
                continue
            elif 'realizar' in texto.lower() and 'ações' in texto.lower():
                variaveis_exigidas = parse_num_words(texto)
                continue
            elif tipo == 'atividades pa':
                pa_complementares_raw = texto
                continue
            elif tipo in ['ou', 'atividade substitutiva'] or texto.lower().startswith('ou conquistar'):
                if not texto.lower().startswith('ou conquistar'):
                    substitutivas_list.append({
                        'excel_row': r_idx,
                        'ds_item': texto,
                        'regras': regras
                    })
                continue
            elif tipo == 'fixas':
                fixas_list.append({
                    'excel_row': r_idx,
                    'tp_acao': 'Fixa',
                    'modalidade': 'Básico',
                    'ds_acao': texto,
                    'regras': regras
                })
            elif tipo == 'fixas ar':
                fixas_list.append({
                    'excel_row': r_idx,
                    'tp_acao': 'Fixa',
                    'modalidade': 'Ar',
                    'ds_acao': texto,
                    'regras': regras
                })
            elif tipo == 'fixas mar':
                fixas_list.append({
                    'excel_row': r_idx,
                    'tp_acao': 'Fixa',
                    'modalidade': 'Mar',
                    'ds_acao': texto,
                    'regras': regras
                })
            elif tipo == 'variaveis':
                variaveis_list.append({
                    'excel_row': r_idx,
                    'tp_acao': 'Variável',
                    'modalidade': 'Básico',
                    'ds_acao': texto,
                    'regras': regras
                })
            elif tipo == 'variaveis ar':
                variaveis_list.append({
                    'excel_row': r_idx,
                    'tp_acao': 'Variável',
                    'modalidade': 'Ar',
                    'ds_acao': texto,
                    'regras': regras
                })
            elif tipo == 'variaveis mar':
                variaveis_list.append({
                    'excel_row': r_idx,
                    'tp_acao': 'Variável',
                    'modalidade': 'Mar',
                    'ds_acao': texto,
                    'regras': regras
                })
            elif 'especialidade' in tipo:
                if not texto.lower().startswith('especialidades que podem') and not texto.lower().startswith('conquistar ao menos uma'):
                    esps_compl.append({
                        'excel_row': r_idx,
                        'nome': texto,
                        'regras': regras
                    })

        b_norm = normalize_text_key(nm_bloco)
        
        # Ação consolidada de especialidades complementares
        if len(esps_compl) > 0:
            is_bloco_1 = ('aprendizagem' in b_norm and 'vocacional' in b_norm)
            nivel_min_esp = 1 if is_bloco_1 else 2
            nomes_esps = [e['nome'] for e in esps_compl]
            ds_esps_acao = f"Conquistar ao menos uma das seguintes especialidades no nível {nivel_min_esp}+: " + ", ".join(nomes_esps)
            variaveis_list.append({
                'excel_row': 0,
                'tp_acao': 'Variável',
                'modalidade': 'Básico',
                'ds_acao': ds_esps_acao,
                'regras': ['especialidade', f"{nivel_min_esp}+"],
                'is_esps_compl': True,
                'esps_list': nomes_esps,
                'nivel_min_esp': nivel_min_esp
            })

        fixas_basicas_obrigatorias = len([f for f in fixas_list if f.get('modalidade') == 'Básico'])
        fixas_obrigatorias = fixas_basicas_obrigatorias if fixas_basicas_obrigatorias > 0 else len(fixas_list)
        
        # 3. Processar Atividades PA Complementares do Bloco
        pt_refs, rt_refs = parse_atividades_pa_bloco(pa_complementares_raw)
        print(f"  Fixas: {len(fixas_list)} (Obrig: {fixas_obrigatorias}) | Variáveis: {len(variaveis_list)} (Exig: {variaveis_exigidas}) | Especialidades: {len(esps_compl)} | PA: {len(pt_refs)} PTs, {len(rt_refs)} RTs | Subst: {len(substitutivas_list)}")
        
        # 4. Criar Bloco estruturado
        bloco_obj = {
            'eixo': eixo_nm,
            'bloco': nm_bloco,
            'nr_acoes_fixas_obrigatorias': fixas_obrigatorias,
            'nr_acoes_variaveis_exigidas': variaveis_exigidas,
            'nr_ordem': bloco_global_ordem,
            'ds_intencionalidade': intencionalidade
        }
        parsed_blocos.append(bloco_obj)
        
        # 5. Processar Fixas
        for fix in fixas_list:
            pistas_set = []
            rumo_set = []
            esp_set = []
            
            regras_str = ' '.join(fix['regras'])
            p_f, r_f = parse_pa_refs_from_text(regras_str)
            pistas_set.extend(p_f)
            rumo_set.extend(r_f)
            
            for rg in fix['regras']:
                e_list, _ = parse_especialidades_from_text(rg)
                for e in e_list:
                    if e not in esp_set: esp_set.append(e)
            
            tp_regra = 'DIRETA'
            total_refs = len(pistas_set) + len(rumo_set) + len(esp_set)
            if total_refs == 0:
                tp_regra = 'SEM_EQUIVALENCIA'
            elif len(esp_set) > 0 and len(pistas_set) == 0 and len(rumo_set) == 0:
                tp_regra = 'ESPECIALIDADES'
            elif total_refs > 1:
                tp_regra = 'OR'
                
            label_fh = ' | '.join(fix['regras']) if fix['regras'] else f"Modalidade do {fix['modalidade']}" if fix['modalidade'] != 'Básico' else 'Alterar manualmente'
            
            parsed_acoes.append({
                'excel_row': fix['excel_row'],
                'eixo': eixo_nm,
                'bloco': nm_bloco,
                'tp_acao': 'Fixa',
                'modalidade': fix['modalidade'],
                'ds_acao': fix['ds_acao'],
                'regra_qtd_texto': ''
            })
            
            parsed_regras.append({
                'eixo': eixo_nm,
                'bloco': nm_bloco,
                'tipo_acao': 'Fixa',
                'excel_row': fix['excel_row'],
                'ds_acao_c': fix['ds_acao'],
                'label_f_h': label_fh,
                'tp_regra': tp_regra,
                'min_count': 1,
                'refs_pistas_ueb': pistas_set,
                'refs_rumo_ueb': rumo_set,
                'refs_especialidades': esp_set
            })
            
        # 6. Processar Variáveis
        for var in variaveis_list:
            pistas_set = []
            rumo_set = []
            esp_set = []
            nivel_min_esp = 1
            
            if var.get('is_esps_compl'):
                esp_set = var['esps_list']
                tp_regra = 'ESPECIALIDADES'
                nivel_min_esp = var.get('nivel_min_esp', 2)
                label_fh = f"Especialidade (Nível {nivel_min_esp}+): {', '.join(esp_set[:6])}" + (f" e mais {len(esp_set)-6} opções" if len(esp_set) > 6 else "")
            else:
                regras_str = ' '.join(var['regras'])
                p_f, r_f = parse_pa_refs_from_text(regras_str)
                pistas_set.extend(p_f)
                rumo_set.extend(r_f)
                
                for rg in var['regras']:
                    e_list, nv = parse_especialidades_from_text(rg)
                    if nv > nivel_min_esp: nivel_min_esp = nv
                    for e in e_list:
                        if e not in esp_set: esp_set.append(e)
                
                tp_regra = 'DIRETA'
                total_refs = len(pistas_set) + len(rumo_set) + len(esp_set)
                if total_refs == 0:
                    tp_regra = 'SEM_EQUIVALENCIA'
                elif len(esp_set) > 0 and len(pistas_set) == 0 and len(rumo_set) == 0:
                    tp_regra = 'ESPECIALIDADES'
                elif total_refs > 1:
                    tp_regra = 'OR'
                label_fh = ' | '.join(var['regras']) if var['regras'] else f"Modalidade do {var['modalidade']}" if var['modalidade'] != 'Básico' else 'Alterar manualmente'
                
            parsed_acoes.append({
                'excel_row': var['excel_row'],
                'eixo': eixo_nm,
                'bloco': nm_bloco,
                'tp_acao': 'Variável',
                'modalidade': var['modalidade'],
                'ds_acao': var['ds_acao'],
                'regra_qtd_texto': f"Realizar {variaveis_exigidas} ações"
            })
            
            parsed_regras.append({
                'eixo': eixo_nm,
                'bloco': nm_bloco,
                'tipo_acao': 'Variável',
                'excel_row': var['excel_row'],
                'ds_acao_c': var['ds_acao'],
                'label_f_h': label_fh,
                'tp_regra': tp_regra,
                'min_count': 1,
                'nivel_min_especialidade': nivel_min_esp,
                'refs_pistas_ueb': pistas_set,
                'refs_rumo_ueb': rumo_set,
                'refs_especialidades': esp_set
            })
            
        # 7. Injetar Atividades PA Complementares na sequência das Variáveis com TAG 'PA'
        for pt_num in pt_refs:
            pa_key = f"PT-{pt_num}"
            pa_info = pa_atividades_map.get(pa_key, {})
            ds_texto = pa_info.get('ds_atividade') or f"Atividade de Pistas e Trilha nº {pt_num}"
            
            parsed_acoes.append({
                'excel_row': 0,
                'eixo': eixo_nm,
                'bloco': nm_bloco,
                'tp_acao': 'Variável',
                'modalidade': 'PA',
                'ds_acao': ds_texto,
                'regra_qtd_texto': f"PA Complementar (Pista {pt_num})"
            })
            
            parsed_regras.append({
                'eixo': eixo_nm,
                'bloco': nm_bloco,
                'tipo_acao': 'Variável',
                'excel_row': 0,
                'ds_acao_c': ds_texto,
                'label_f_h': f"Pista {pt_num}",
                'tp_regra': 'DIRETA',
                'min_count': 1,
                'refs_pistas_ueb': [pt_num],
                'refs_rumo_ueb': [],
                'refs_especialidades': []
            })
            
        for rt_num in rt_refs:
            pa_key = f"RT-{rt_num}"
            pa_info = pa_atividades_map.get(pa_key, {})
            ds_texto = pa_info.get('ds_atividade') or f"Atividade de Rumo e Travessia nº {rt_num}"
            
            parsed_acoes.append({
                'excel_row': 0,
                'eixo': eixo_nm,
                'bloco': nm_bloco,
                'tp_acao': 'Variável',
                'modalidade': 'PA',
                'ds_acao': ds_texto,
                'regra_qtd_texto': f"PA Complementar (Rumo {rt_num})"
            })
            
            parsed_regras.append({
                'eixo': eixo_nm,
                'bloco': nm_bloco,
                'tipo_acao': 'Variável',
                'excel_row': 0,
                'ds_acao_c': ds_texto,
                'label_f_h': f"Rumo {rt_num}",
                'tp_regra': 'DIRETA',
                'min_count': 1,
                'refs_pistas_ueb': [],
                'refs_rumo_ueb': [rt_num],
                'refs_especialidades': []
            })
            
        # 8. Injetar Atividades Substitutivas
        for sub in substitutivas_list:
            ds_sub = sub['ds_item']
            parsed_acoes.append({
                'excel_row': sub['excel_row'],
                'eixo': eixo_nm,
                'bloco': nm_bloco,
                'tp_acao': 'Substitutiva',
                'modalidade': 'Substitutiva',
                'ds_acao': ds_sub,
                'regra_qtd_texto': 'Atividade Substitutiva do Bloco'
            })
            
            parsed_regras.append({
                'eixo': eixo_nm,
                'bloco': nm_bloco,
                'tipo_acao': 'Substitutiva',
                'excel_row': sub['excel_row'],
                'ds_acao_c': ds_sub,
                'label_f_h': f"Especialidade/Insígnia: {ds_sub}",
                'tp_regra': 'ESPECIALIDADES',
                'min_count': 1,
                'refs_pistas_ueb': [],
                'refs_rumo_ueb': [],
                'refs_especialidades': [ds_sub]
            })

print(f"\n=======================================================")
print(f"★ RESUMO DO PROCESSAMENTO ★")
print(f"Total de Blocos: {len(parsed_blocos)}")
print(f"Total de Ações PN (Fixas + Variáveis + PA + Substitutivas): {len(parsed_acoes)}")
print(f"Total de Regras de Equivalência: {len(parsed_regras)}")
print(f"=======================================================")

pn_out = {
    'eixos': eixos_list,
    'blocos': parsed_blocos,
    'acoes': parsed_acoes
}

equiv_out = {
    'total_regras': len(parsed_regras),
    'regras': parsed_regras
}

with open(os.path.join(BASE_DIR, 'data/catalogo/pn_catalogo.json'), 'w', encoding='utf-8') as f:
    json.dump(pn_out, f, ensure_ascii=False, indent=2)

with open(os.path.join(BASE_DIR, 'data/catalogo/equivalencias.json'), 'w', encoding='utf-8') as f:
    json.dump(equiv_out, f, ensure_ascii=False, indent=2)

print("✓ Arquivos pn_catalogo.json e equivalencias.json salvos com sucesso!")
